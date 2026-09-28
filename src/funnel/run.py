"""Weekly raw-to-report pipeline. Run from the repository root."""

from __future__ import annotations

import argparse
import base64
import hashlib
import logging
from pathlib import Path
import sys
import time
from datetime import date, datetime

import pandas as pd
from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .analysis import cohort_table, grouped_rates, plot_bar, plot_cohorts
from .cleaning import DATE_COLUMNS, NUMERIC_COLUMNS, clean
from .schema import SchemaError, mapped_groups, validate


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_INPUT = "Case_Bari/propostas_credito.csv"
DEFAULT_OUTPUT = "outputs/"
MAX_PARSE_FAILURE_RATE = 0.02
TEMPLATE_DIR = Path(__file__).parent / "templates"


def resolve_root_path(value: str) -> Path:
    path = Path(value)
    return path if path.is_absolute() else ROOT / path


def display_path(path: Path) -> str:
    try:
        return str(path.relative_to(ROOT))
    except ValueError:
        return str(path)


def setup_logger(output_dir: Path) -> tuple[logging.Logger, Path]:
    log_dir = output_dir / "logs"
    log_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    path = log_dir / f"run_{stamp}.log"
    logger = logging.getLogger(f"funnel.weekly.{stamp}")
    logger.setLevel(logging.INFO)
    logger.propagate = False
    for handler in (logging.FileHandler(path, encoding="utf-8"), logging.StreamHandler(sys.stdout)):
        handler.setFormatter(logging.Formatter("%(levelname)s %(message)s"))
        logger.addHandler(handler)
    return logger, path


def parse_quality(cleaned: pd.DataFrame, maximum_rate: float) -> tuple[int, int]:
    columns = NUMERIC_COLUMNS + DATE_COLUMNS
    failures = sum(int(cleaned[f"flag_parse_error_{column}"].sum()) for column in columns)
    cells = sum(int(cleaned[f"{column}_raw"].astype("string").str.strip().replace("", pd.NA).notna().sum())
                for column in columns)
    if cells and failures / cells > maximum_rate:
        raise SchemaError(f"Falhas de conversão: {failures}/{cells} células ({failures/cells:.2%}); "
                          f"limiar configurado {maximum_rate:.2%}. Corrija a origem ou revise o limiar.")
    return failures, cells


def data_tables(df: pd.DataFrame, as_of: pd.Timestamp, quality: pd.DataFrame,
                warnings: list[str]) -> dict[str, pd.DataFrame | int | str]:
    n = len(df)
    contracts = int(df.flag_contratada.sum())
    available = df.tempo_analise_dias.dropna()
    horizon = int(available.quantile(0.9).__ceil__()) if len(available) else None
    mature = ((as_of - df.data_entrada).dt.days.ge(horizon) & df.data_entrada.notna()) if horizon else pd.Series(False, index=df.index)
    dates = df.loc[df.data_entrada.notna()].copy()
    trend = cohort_table(dates, "M", mature.loc[dates.index]) if len(dates) else pd.DataFrame(
        columns=["coorte", "N total", "contratos total", "taxa ingênua", "N maduro", "contratos maduros", "taxa madura"])
    # Group unknown status explicitly, without changing cleaned source categories.
    status_groups = mapped_groups(df, "status_final")
    stage_groups = df.etapa_max_funil.astype("string").fillna("Indisponível")
    loss = df.loc[~df.flag_contratada].copy()
    loss["status_group"] = status_groups.loc[loss.index]
    loss["stage_group"] = stage_groups.loc[loss.index]
    funnel = loss.groupby(["stage_group", "status_group"], dropna=False).agg(
        Propostas=("id_proposta", "size"), **{"Crédito solicitado (R$)": ("valor_solicitado", "sum")}
    ).reset_index().rename(columns={"stage_group": "Etapa máxima", "status_group": "Status"})
    funnel["Ticket médio (R$)"] = funnel["Crédito solicitado (R$)"] / funnel["Propostas"]
    channel = grouped_rates(df, mapped_groups(df, "canal_origem"), 0.95)
    channel.columns = ["Canal", "Propostas", "Contratos", "Conversão", "IC inferior", "IC superior"]
    quality_rows = [{"Indicador": "Linhas analisadas", "Quantidade": n},
                    {"Indicador": "Contratos", "Quantidade": contracts},
                    {"Indicador": "Maturidade estimada (dias, p90)", "Quantidade": horizon if horizon else "indisponível"}]
    for _, row in quality.iterrows():
        quality_rows.append({"Indicador": f"{row['problema']} — {row['coluna']}",
                             "Quantidade": int(row["linhas afetadas"])})
    quality_rows.extend({"Indicador": "Aviso", "Quantidade": warning} for warning in warnings)
    return {"summary": pd.DataFrame([
        {"Indicador": "Propostas", "Valor": n}, {"Indicador": "Contratos", "Valor": contracts},
        {"Indicador": "Conversão", "Valor": contracts / n if n else None},
        {"Indicador": "Crédito solicitado sem contratação (R$)",
         "Valor": float(df.loc[~df.flag_contratada, "valor_solicitado"].sum())},
        {"Indicador": "Horizonte de maturidade estimado (dias)", "Valor": horizon if horizon else "indisponível"},
    ]), "funnel": funnel, "channel": channel, "trend": trend,
        "quality": pd.DataFrame(quality_rows), "mature_count": int(mature.sum())}


def render_charts(tables: dict, output: Path) -> dict[str, str]:
    figure_dir = output / "figures"
    figure_dir.mkdir(parents=True, exist_ok=True)
    funnel = tables["funnel"]
    by_stage = funnel.groupby("Etapa máxima", dropna=False)["Crédito solicitado (R$)"].sum().reset_index()
    by_stage["R$ milhões"] = by_stage["Crédito solicitado (R$)"] / 1_000_000
    paths = {"funnel": figure_dir / "semanal_funil.png", "trend": figure_dir / "semanal_tendencia.png"}
    plot_bar(by_stage, "Etapa máxima", "R$ milhões", "Crédito solicitado sem contratação por etapa",
             "Crédito solicitado (R$ milhões)", paths["funnel"])
    plot_cohorts(tables["trend"].rename(columns={"coorte": "coorte"}), "Contratação por mês de entrada",
                 paths["trend"])
    return {name: base64.b64encode(path.read_bytes()).decode("ascii") for name, path in paths.items()}


def export_html(tables: dict, metadata: dict, figures: dict[str, str], output: Path) -> None:
    def display(value: object, column: str, indicator: str = "") -> str:
        if pd.isna(value):
            return "indisponível"
        if isinstance(value, (float, int)) and not isinstance(value, bool):
            percent = column in {"Conversão", "IC inferior", "IC superior", "taxa ingênua", "taxa madura"}
            percent = percent or indicator == "Conversão"
            currency = "(R$)" in column or "(R$)" in indicator
            if percent:
                return f"{float(value) * 100:,.2f}%".replace(",", "_").replace(".", ",").replace("_", ".")
            if currency:
                return "R$ " + f"{float(value):,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")
            if isinstance(value, float) and not value.is_integer():
                return f"{value:.2f}".replace(".", ",")
        return str(value)

    formatted = {}
    for key, frame in tables.items():
        if isinstance(frame, pd.DataFrame):
            formatted[key] = [{column: display(value, column, str(row.get("Indicador", "")))
                               for column, value in row.items()} for row in frame.to_dict("records")]
    env = Environment(loader=FileSystemLoader(TEMPLATE_DIR), autoescape=select_autoescape(["html"]))
    template = env.get_template("weekly.html.j2")
    html = template.render(meta=metadata, tables=formatted, figures=figures)
    (output / "relatorio_semanal.html").write_text(html, encoding="utf-8")


def export_excel(tables: dict, metadata: dict, output: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = [("Resumo", "summary"), ("Funil", "funnel"), ("Por canal", "channel"),
              ("Tendência", "trend"), ("Qualidade dos dados", "quality")]
    for name, key in sheets:
        ws = workbook.create_sheet(name)
        ws.append([f"Relatório semanal — {name}"])
        for k, value in metadata.items():
            ws.append([k, str(value)])
        ws.append([])
        frame = tables[key]
        ws.append(list(frame.columns))
        header_row = ws.max_row
        for values in frame.itertuples(index=False, name=None):
            ws.append([None if pd.isna(value) else (value.item() if hasattr(value, "item") else value)
                       for value in values])
        for row in ws.iter_rows(min_row=header_row + 1):
            for cell in row:
                column_name = str(ws.cell(header_row, cell.column).value)
                indicator = str(ws.cell(cell.row, 1).value) if key == "summary" else ""
                if column_name in {"Conversão", "IC inferior", "IC superior", "taxa ingênua", "taxa madura"} or indicator == "Conversão":
                    cell.number_format = "0.00%"
                elif "(R$)" in column_name or "(R$)" in indicator:
                    cell.number_format = '"R$ "#,##0.00'
        ws.freeze_panes = f"A{header_row + 1}"
        ws.auto_filter.ref = f"A{header_row}:{ws.cell(ws.max_row, frame.shape[1]).coordinate}"
        for cell in ws[header_row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="265C83")
        for col in ws.columns:
            letter = col[0].column_letter
            width = min(65, max(15, *(len(str(cell.value or "")) + 2 for cell in col[:100])))
            ws.column_dimensions[letter].width = width
    workbook.save(output / "relatorio_semanal.xlsx")


def run(input_path: Path, output_dir: Path, as_of: date | None = None,
        exclude_types: list[str] | None = None, max_parse_failure_rate: float = MAX_PARSE_FAILURE_RATE) -> dict:
    started = time.perf_counter()
    logger, log_path = setup_logger(output_dir)
    try:
        if not 0 <= max_parse_failure_rate <= 1:
            raise SchemaError("O limiar de falhas deve estar entre 0 e 1.")
        logger.info("Início; origem=%s", display_path(input_path))
        if not input_path.is_file():
            raise SchemaError(f"Arquivo de entrada inexistente: {display_path(input_path)}")
        digest = hashlib.sha256(input_path.read_bytes()).hexdigest()
        logger.info("SHA-256=%s", digest)
        raw = pd.read_csv(input_path, dtype="string", keep_default_na=False, encoding="utf-8-sig")
        logger.info("Linhas lidas=%d; validação de schema", len(raw))
        checked = validate(raw)
        for warning in checked.warnings:
            logger.warning(warning)
        logger.info("Tratamento")
        cleaned, treatment = clean(checked.frame, {"EXCLUDE_PROPERTY_TYPES": exclude_types or []})
        warnings = list(checked.warnings)
        for _, row in treatment.iterrows():
            logger.info("Tratamento: %s | %s | %d | %s", row["problema"], row["coluna"],
                        row["linhas afetadas"], row["ação"])
        failures, cells = parse_quality(cleaned, max_parse_failure_rate)
        if failures:
            warning = f"Conversões inválidas: {failures}/{cells} células; valores ausentes nas métricas dependentes."
            warnings.append(warning)
            logger.warning(warning)
        if as_of is None:
            available = cleaned.data_entrada + pd.to_timedelta(cleaned.tempo_analise_dias, unit="D")
            if available.notna().sum() == 0:
                raise SchemaError("Não é possível estimar a data de referência; informe --as-of.")
            cutoff = available.max().normalize()
            warnings.append("Data de referência inferida por entrada + tempo de análise; maturidade é aproximação.")
        else:
            cutoff = pd.Timestamp(as_of)
        future = cleaned.data_entrada.gt(cutoff).fillna(False)
        if future.any():
            warning = f"Entradas posteriores à data de referência: {int(future.sum())}; excluídas deste relatório."
            warnings.append(warning)
            logger.warning(warning)
            cleaned = cleaned.loc[~future].copy()
        if cleaned.empty:
            raise SchemaError("Nenhuma proposta elegível para a data de referência.")
        removed = len(raw) - len(cleaned)
        logger.info("Métricas; referência=%s; linhas analisadas=%d; removidas=%d", cutoff.date(), len(cleaned), removed)
        tables = data_tables(cleaned, cutoff, treatment, warnings)
        valid_basic = int((cleaned.data_entrada.notna() & cleaned.valor_solicitado.notna()
                           & ~cleaned.etapa_invalida & cleaned.etapa_max_funil.notna()).sum())
        tables["summary"] = pd.concat([tables["summary"], pd.DataFrame([{
            "Indicador": "Coerência de taxa e contratação",
            "Valor": "indisponível" if checked.missing_optional else int(cleaned.flag_status_taxa_incoerente.sum()),
        }])], ignore_index=True)
        metadata = {
            "Gerado em": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "Arquivo": display_path(input_path), "SHA-256": digest, "Data de referência": cutoff.date().isoformat(),
            "Linhas lidas": len(raw), "Linhas analisadas": len(cleaned),
            "Linhas com campos básicos válidos": valid_basic, "Linhas removidas": removed,
            "Avisos": " | ".join(warnings) if warnings else "Nenhum",
        }
        logger.info("Exportação HTML, Excel e CSV tratado")
        figures = render_charts(tables, output_dir)
        export_html(tables, metadata, figures, output_dir)
        export_excel(tables, metadata, output_dir)
        cleaned.to_csv(output_dir / "propostas_weekly_clean.csv", index=False)
        logger.info("Concluído em %.2fs; HTML=%s; Excel=%s; log=%s",
                    time.perf_counter() - started, display_path(output_dir / "relatorio_semanal.html"),
                    display_path(output_dir / "relatorio_semanal.xlsx"), display_path(log_path))
        return {"metadata": metadata, "tables": tables, "log_path": log_path}
    except (SchemaError, ValueError, pd.errors.ParserError, UnicodeError, OSError) as exc:
        logger.error("Execução interrompida: %s", exc)
        raise
    except Exception as exc:
        logger.error("Falha inesperada na rotina: %s: %s", type(exc).__name__, exc)
        raise
    finally:
        for handler in logger.handlers[:]:
            handler.close()
            logger.removeHandler(handler)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Gera o relatório semanal do funil a partir do CSV bruto.")
    parser.add_argument("--input", default=DEFAULT_INPUT)
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT)
    parser.add_argument("--as-of", type=date.fromisoformat, metavar="YYYY-MM-DD")
    parser.add_argument("--exclude-types", default="", help="Lista separada por vírgulas; vazia por padrão.")
    parser.add_argument("--max-parse-failure-rate", type=float, default=MAX_PARSE_FAILURE_RATE,
                        help="Fração máxima de células não vazias que falharam na conversão.")
    args = parser.parse_args(argv)
    try:
        run(resolve_root_path(args.input), resolve_root_path(args.output_dir), args.as_of,
            [value.strip() for value in args.exclude_types.split(",") if value.strip()],
            args.max_parse_failure_rate)
        return 0
    except Exception as exc:
        if isinstance(exc, (KeyboardInterrupt, SystemExit)):
            raise
        print(f"Falha no relatório semanal: {exc}. Consulte o log em {args.output_dir}/logs/.", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
